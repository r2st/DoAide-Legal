import io
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from generators import format_inr


BRAND_FOOTER = "Generated with DoAide Legal — legal.doaide.com"


def _create_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
    return doc


def _add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    return p


def _add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(128, 128, 128)
    return p


def _add_section_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    return p


def _add_body(doc, text, bold=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run(text)
    run.font.size = Pt(11)
    if bold:
        run.bold = True
    return p


def _add_footer(doc):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(BRAND_FOOTER)
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(153, 153, 153)


def _to_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _ordinal(n: int) -> str:
    if 11 <= (n % 100) <= 13:
        suffix = "th"
    else:
        suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def generate_rental_agreement(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, "RENTAL / LEASE AGREEMENT")
    _add_body(doc, f"This Rental Agreement is made and executed on {data.lease_start_date} at the place mentioned herein between the parties described below:")
    doc.add_paragraph()

    _add_body(doc, "LANDLORD / LESSOR:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    text = f"Name: {data.landlord_name}\nAddress: {data.landlord_address}"
    if data.landlord_aadhaar:
        text += f"\nAadhaar: {data.landlord_aadhaar}"
    _add_body(doc, text, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, '(Hereinafter referred to as the "Landlord", which expression shall mean and include their heirs, successors, legal representatives, and assigns)')

    _add_body(doc, "TENANT / LESSEE:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    text = f"Name: {data.tenant_name}\nAddress: {data.tenant_address}"
    if data.tenant_aadhaar:
        text += f"\nAadhaar: {data.tenant_aadhaar}"
    _add_body(doc, text, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, '(Hereinafter referred to as the "Tenant", which expression shall mean and include their heirs, successors, legal representatives, and assigns)')

    _add_section_heading(doc, "PROPERTY DETAILS")
    _add_body(doc, f"Address: {data.property_address}\nType: {data.property_type.title()}\nFurnishing: {data.furnishing.replace('-', ' ').title()}\nPurpose: {data.purpose.title()}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "TERMS AND CONDITIONS")
    clauses = [
        f"1. Rent: The Tenant shall pay a monthly rent of {format_inr(data.monthly_rent)} to the Landlord on or before the {_ordinal(data.rent_due_day)} of each calendar month.",
        f"2. Security Deposit: The Tenant has paid a security deposit of {format_inr(data.security_deposit)} to the Landlord, which shall be refunded at the time of vacating the premises after deducting any dues or damages.",
        f"3. Duration: This agreement shall be valid for a period of {data.lease_duration_months} months commencing from {data.lease_start_date}.",
        f"4. Notice Period: Either party may terminate this agreement by giving {data.notice_period_months} month(s) written notice to the other party.",
        "5. Usage: The Tenant shall use the premises solely for the purpose mentioned above.",
        "6. Sub-letting: The Tenant shall not sub-let, assign, or transfer the tenancy without prior written consent of the Landlord.",
        "7. Maintenance: The Tenant shall maintain the premises in good condition. Major structural repairs shall be the Landlord's responsibility.",
        "8. Alterations: The Tenant shall not make structural alterations without the Landlord's written consent.",
        "9. Inspection: The Landlord may inspect the premises at reasonable hours with prior notice.",
        "10. Utilities: The Tenant shall bear all charges for electricity, water, gas, and other utilities.",
        "11. Vacating: Upon expiry or termination, the Tenant shall vacate and hand over peaceful possession.",
        "12. Disputes: Disputes shall be subject to the jurisdiction of the courts at the property's location.",
    ]
    if data.maintenance_charges:
        clauses.append(f"13. Maintenance Charges: The Tenant shall pay maintenance charges of {format_inr(data.maintenance_charges)} per month.")
    for c in clauses:
        _add_body(doc, c)
    if data.additional_clauses:
        start = len(clauses) + 1
        for i, clause in enumerate(data.additional_clauses):
            _add_body(doc, f"{start + i}. {clause}")

    doc.add_paragraph()
    _add_body(doc, "IN WITNESS WHEREOF, the parties have set their hands on the day and year first above written.")
    doc.add_paragraph()
    _add_body(doc, f"LANDLORD: {data.landlord_name}\nSignature: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, f"TENANT: {data.tenant_name}\nSignature: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, "WITNESSES:\n1. Name: _______________ Signature: _______________\n2. Name: _______________ Signature: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_nda(data) -> bytes:
    doc = _create_doc()
    nda_label = "MUTUAL" if data.nda_type == "mutual" else "ONE-WAY"
    _add_title(doc, f"{nda_label} NON-DISCLOSURE AGREEMENT")

    _add_body(doc, f'This Non-Disclosure Agreement ("Agreement") is entered into on {data.effective_date} by and between:')
    if data.nda_type == "mutual":
        _add_body(doc, f"Party A: {data.disclosing_party_name}, {data.disclosing_party_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
        _add_body(doc, f"Party B: {data.receiving_party_name}, {data.receiving_party_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    else:
        _add_body(doc, f"Disclosing Party: {data.disclosing_party_name}, {data.disclosing_party_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
        _add_body(doc, f"Receiving Party: {data.receiving_party_name}, {data.receiving_party_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "RECITALS")
    _add_body(doc, f"WHEREAS, the parties wish to explore a potential business relationship regarding: {data.purpose}, and confidential information may be disclosed.")

    _add_section_heading(doc, "DEFINITIONS")
    _add_body(doc, '"Confidential Information" means any information disclosed by one party to the other, including trade secrets, business plans, financial data, customer lists, and technical specifications.')

    _add_section_heading(doc, "OBLIGATIONS")
    for o in ["1. Hold all Confidential Information in strict confidence.", "2. Use it solely for the stated purpose.", "3. Protect it with at least reasonable care.", "4. Limit access to employees and advisors with a need to know."]:
        _add_body(doc, o)

    _add_section_heading(doc, "EXCLUSIONS")
    for ex in ["1. Publicly available information.", "2. Already known before disclosure.", "3. Independently developed.", "4. Disclosed pursuant to court order."]:
        _add_body(doc, ex)
    if data.additional_exclusions:
        for i, ex in enumerate(data.additional_exclusions):
            _add_body(doc, f"{5 + i}. {ex}")

    _add_section_heading(doc, "TERM")
    _add_body(doc, f"This Agreement remains in effect for {data.confidentiality_period_years} year(s).")

    _add_section_heading(doc, "GOVERNING LAW")
    _add_body(doc, f"Governed by Indian law; courts of {data.governing_state} shall have exclusive jurisdiction.")

    doc.add_paragraph()
    _add_body(doc, f"Disclosing Party: {data.disclosing_party_name}\nSignature: _______________\nDate: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, f"Receiving Party: {data.receiving_party_name}\nSignature: _______________\nDate: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_offer_letter(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, data.company_name)
    _add_subtitle(doc, data.company_address)

    _add_body(doc, f"Date: {data.offer_date}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_body(doc, f"To,\n{data.candidate_name}\n{data.candidate_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"Subject: Offer of Employment — {data.designation}", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"Dear {data.candidate_name},")
    _add_body(doc, f"We are pleased to offer you the position of {data.designation} in the {data.department} department at {data.company_name}. Your expected date of joining is {data.date_of_joining}.")

    _add_section_heading(doc, "COMPENSATION DETAILS")
    table = doc.add_table(rows=6, cols=3)
    table.style = "Table Grid"
    headers = ["Component", "Monthly (₹)", "Annual (₹)"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    rows_data = [
        ("Basic Salary", data.basic_salary, data.basic_salary * 12),
        ("HRA", data.hra, data.hra * 12),
        ("Special Allowance", data.special_allowance, data.special_allowance * 12),
        ("PF (Employer)", data.pf_contribution, data.pf_contribution * 12),
        ("CTC", "", data.ctc_annual),
    ]
    for i, (name, monthly, annual) in enumerate(rows_data):
        table.rows[i + 1].cells[0].text = name
        table.rows[i + 1].cells[1].text = format_inr(monthly) if isinstance(monthly, (int, float)) else str(monthly)
        table.rows[i + 1].cells[2].text = format_inr(annual)

    doc.add_paragraph()
    _add_section_heading(doc, "OTHER TERMS")
    for t in [f"Reporting Manager: {data.reporting_manager}", f"Work Location: {data.work_location}", f"Probation: {data.probation_months} months", f"Notice Period: {data.notice_period_months} month(s)"]:
        _add_body(doc, f"• {t}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if data.additional_benefits:
        _add_body(doc, "Additional Benefits:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        for b in data.additional_benefits:
            _add_body(doc, f"• {b}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_body(doc, f"Please confirm acceptance by {data.offer_expiry_date}.")
    doc.add_paragraph()
    _add_body(doc, f"For {data.company_name}\n_______________\nAuthorized Signatory", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, "ACCEPTANCE", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"I, {data.candidate_name}, hereby accept the above offer.\nSignature: _______________ Date: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_freelancer_contract(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, "FREELANCER SERVICE AGREEMENT")
    _add_body(doc, f'This Service Agreement is entered into on {data.start_date} between:')
    _add_body(doc, f"Client: {data.client_name}, {data.client_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"Freelancer: {data.freelancer_name}, {data.freelancer_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "1. SCOPE OF WORK")
    _add_body(doc, data.project_description)
    _add_body(doc, "Deliverables:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    for i, d in enumerate(data.deliverables):
        _add_body(doc, f"{i + 1}. {d}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "2. TIMELINE")
    _add_body(doc, f"Start: {data.start_date} | End: {data.end_date}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "3. COMPENSATION")
    _add_body(doc, f"Total Fee: {format_inr(data.total_fee)} | Payment: {data.payment_schedule.title()}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if data.advance_percentage:
        _add_body(doc, f"Advance: {data.advance_percentage}% ({format_inr(data.total_fee * data.advance_percentage / 100)})", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "4. REVISIONS")
    _add_body(doc, f"Client entitled to {data.revision_rounds} revision rounds. Additional revisions at extra cost.")

    _add_section_heading(doc, "5. INTELLECTUAL PROPERTY")
    ip_map = {"client": "All IP vests with the Client.", "freelancer": "Freelancer retains IP; Client gets perpetual license.", "shared": "Joint ownership."}
    _add_body(doc, ip_map.get(data.ip_ownership, ip_map["client"]))

    if data.confidentiality:
        _add_section_heading(doc, "6. CONFIDENTIALITY")
        _add_body(doc, "Both parties agree to keep project information confidential.")

    _add_section_heading(doc, "7. TERMINATION")
    _add_body(doc, f"Either party may terminate with {data.termination_notice_days} days' notice. Freelancer compensated for completed work.")

    _add_section_heading(doc, "8. GOVERNING LAW")
    _add_body(doc, f"Governed by Indian law; jurisdiction in {data.governing_state}.")

    doc.add_paragraph()
    _add_body(doc, f"Client: {data.client_name}\nSignature: _______________\nDate: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, f"Freelancer: {data.freelancer_name}\nSignature: _______________\nDate: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_invoice(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, "TAX INVOICE")

    _add_body(doc, data.seller_name, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, data.seller_address, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if data.seller_gstin:
        _add_body(doc, f"GSTIN: {data.seller_gstin}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if data.seller_pan:
        _add_body(doc, f"PAN: {data.seller_pan}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_body(doc, f"Invoice #: {data.invoice_number} | Date: {data.invoice_date} | Due: {data.due_date}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()

    _add_body(doc, "Bill To:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"{data.buyer_name}\n{data.buyer_address}" + (f"\nGSTIN: {data.buyer_gstin}" if data.buyer_gstin else ""), alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()

    has_hsn = any(item.hsn_code for item in data.items)
    cols = 5 + (1 if has_hsn else 0)
    table = doc.add_table(rows=1, cols=cols)
    table.style = "Table Grid"
    headers = ["#", "Description"]
    if has_hsn:
        headers.append("HSN")
    headers.extend(["Qty", "Rate (₹)", "Amount (₹)"])
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    subtotal = 0.0
    for i, item in enumerate(data.items):
        amount = item.quantity * item.rate
        subtotal += amount
        row = table.add_row()
        vals = [str(i + 1), item.description]
        if has_hsn:
            vals.append(item.hsn_code or "")
        vals.extend([str(item.quantity), format_inr(item.rate), format_inr(amount)])
        for j, v in enumerate(vals):
            row.cells[j].text = v

    doc.add_paragraph()
    discount = subtotal * (data.discount_percentage or 0) / 100.0
    taxable = subtotal - discount
    _add_body(doc, f"Subtotal: {format_inr(subtotal)}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    if discount > 0:
        _add_body(doc, f"Discount ({data.discount_percentage}%): -{format_inr(discount)}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_body(doc, f"Taxable: {format_inr(taxable)}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    if data.is_igst:
        igst = taxable * data.gst_rate / 100.0
        _add_body(doc, f"IGST ({data.gst_rate}%): {format_inr(igst)}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        grand_total = taxable + igst
    else:
        half = data.gst_rate / 2.0
        cgst = taxable * half / 100.0
        _add_body(doc, f"CGST ({half}%): {format_inr(cgst)}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        _add_body(doc, f"SGST ({half}%): {format_inr(cgst)}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        grand_total = taxable + cgst * 2
    _add_body(doc, f"GRAND TOTAL: {format_inr(grand_total)}", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)

    if data.bank_name:
        _add_section_heading(doc, "BANK DETAILS")
        bank_text = f"Bank: {data.bank_name}"
        if data.account_number:
            bank_text += f"\nA/C No: {data.account_number}"
        if data.ifsc_code:
            bank_text += f"\nIFSC: {data.ifsc_code}"
        _add_body(doc, bank_text, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    if data.notes:
        _add_body(doc, f"Notes: {data.notes}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_paragraph()
    _add_body(doc, f"For {data.seller_name}\nAuthorized Signatory", alignment=WD_ALIGN_PARAGRAPH.RIGHT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_power_of_attorney(data) -> bytes:
    doc = _create_doc()
    poa_label = "GENERAL" if data.poa_type == "general" else "SPECIAL"
    _add_title(doc, f"{poa_label} POWER OF ATTORNEY")

    principal_text = f"KNOW ALL MEN BY THESE PRESENTS that I, {data.principal_name}, residing at {data.principal_address}"
    if data.principal_aadhaar:
        principal_text += f" (Aadhaar: {data.principal_aadhaar})"
    principal_text += f", do hereby appoint {data.attorney_name}, residing at {data.attorney_address}"
    if data.attorney_aadhaar:
        principal_text += f" (Aadhaar: {data.attorney_aadhaar})"
    principal_text += " as my true and lawful Attorney."
    _add_body(doc, principal_text)

    _add_section_heading(doc, "POWERS GRANTED")
    for i, p in enumerate(data.powers_granted):
        _add_body(doc, f"{i + 1}. {p}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "TERMS")
    term_text = f"Effective from {data.effective_date}"
    term_text += f", valid until {data.expiry_date}." if data.expiry_date else ", until revoked in writing."
    _add_body(doc, term_text)
    _add_body(doc, f"Governed by Indian law; jurisdiction in {data.governing_state}.")

    doc.add_paragraph()
    _add_body(doc, f"PRINCIPAL: {data.principal_name}\nSignature: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, f"ATTORNEY: {data.attorney_name}\nSignature: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, "WITNESSES:", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    for i, w in enumerate(data.witnesses[:2]):
        _add_body(doc, f"{i + 1}. {w.name}, {w.address}\nSignature: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_partnership_deed(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, "PARTNERSHIP DEED")
    _add_body(doc, f"This Partnership Deed is made on {data.commencement_date} by and between:")
    for i, p in enumerate(data.partners):
        _add_body(doc, f"Partner {i+1}: {p.name}, {p.address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "1. NAME OF FIRM")
    _add_body(doc, f'The firm shall be known as "{data.firm_name}".')
    _add_section_heading(doc, "2. PLACE OF BUSINESS")
    _add_body(doc, data.firm_address)
    _add_section_heading(doc, "3. NATURE OF BUSINESS")
    _add_body(doc, data.business_nature)
    _add_section_heading(doc, "4. COMMENCEMENT")
    duration_text = "at the will of the partners" if data.duration == "at-will" else "for a fixed term"
    _add_body(doc, f"Commencing {data.commencement_date}, {duration_text}.")

    _add_section_heading(doc, "5. CAPITAL & PROFIT SHARING")
    table = doc.add_table(rows=1 + len(data.partners), cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Partner", "Capital (₹)", "Profit (%)"]):
        table.rows[0].cells[i].text = h
    for i, p in enumerate(data.partners):
        table.rows[i + 1].cells[0].text = p.name
        table.rows[i + 1].cells[1].text = format_inr(p.capital_contribution)
        table.rows[i + 1].cells[2].text = f"{p.profit_share_percentage}%"

    _add_section_heading(doc, "6. BANKING")
    _add_body(doc, f"Bank: {data.bank_name}" + (f", A/C: {data.bank_account}" if data.bank_account else ""), alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_section_heading(doc, "7. FINANCIAL YEAR")
    _add_body(doc, f"Commences {data.financial_year_start} each year.")
    _add_section_heading(doc, "8. DISPUTE RESOLUTION")
    _add_body(doc, f"{'Arbitration under A&C Act, 1996' if data.dispute_resolution == 'arbitration' else 'Courts'} in {data.governing_state}.")

    if data.additional_clauses:
        _add_section_heading(doc, "9. ADDITIONAL CLAUSES")
        for i, c in enumerate(data.additional_clauses):
            _add_body(doc, f"{i + 1}. {c}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_paragraph()
    for i, p in enumerate(data.partners):
        _add_body(doc, f"Partner {i+1}: {p.name}\nSignature: _______________", alignment=WD_ALIGN_PARAGRAPH.LEFT)
        doc.add_paragraph()

    _add_footer(doc)
    return _to_bytes(doc)


def generate_resignation_letter(data) -> bytes:
    doc = _create_doc()
    _add_body(doc, f"Date: {data.resignation_date}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_body(doc, f"To,\n{data.manager_name}\n{data.manager_designation}\n{data.company_name}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, "Subject: Resignation from Services", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"Dear {data.manager_name},")

    body = f"I, {data.employee_name}, working as {data.employee_designation} in the {data.department} department"
    if data.employee_id:
        body += f" (Employee ID: {data.employee_id})"
    body += f", hereby tender my resignation from {data.company_name}."
    _add_body(doc, body)

    if data.reason:
        _add_body(doc, f"Reason: {data.reason}")
    _add_body(doc, f"As per my notice period of {data.notice_period_days} days, my last working day will be {data.last_working_date}.")
    _add_body(doc, "I am willing to complete pending tasks and assist in transition.")
    _add_body(doc, f"I am grateful for my tenure at {data.company_name} and wish the organization success.")
    _add_body(doc, "Thanking you,")
    doc.add_paragraph()
    _add_body(doc, f"Yours sincerely,\n{data.employee_name}\n{data.employee_designation}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_experience_certificate(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, data.company_name)
    _add_subtitle(doc, data.company_address)
    _add_body(doc, f"Date: {data.issue_date}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_title(doc, "EXPERIENCE CERTIFICATE")
    _add_body(doc, "TO WHOM IT MAY CONCERN", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    _add_body(doc, f"This is to certify that {data.employee_name} was employed with {data.company_name} as {data.employee_designation} in the {data.department} department from {data.date_of_joining} to {data.date_of_leaving}.")
    if data.responsibilities:
        _add_body(doc, "During their tenure, they were responsible for:")
        for r in data.responsibilities:
            _add_body(doc, f"• {r}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    if data.performance_rating:
        ratings = {"excellent": "excellent", "good": "good and satisfactory", "satisfactory": "satisfactory"}
        _add_body(doc, f"Their performance was {ratings.get(data.performance_rating, '')}.")
    _add_body(doc, f"We wish {data.employee_name} all the best in their future endeavours.")

    doc.add_paragraph()
    _add_body(doc, f"For {data.company_name}\n\n_______________\n{data.signatory_name}\n{data.signatory_designation}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_salary_slip(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, data.company_name)
    _add_subtitle(doc, data.company_address)
    _add_title(doc, f"SALARY SLIP — {data.month} {data.year}")

    info_lines = [
        f"Employee: {data.employee_name} | ID: {data.employee_id}",
        f"Designation: {data.designation} | Department: {data.department}",
        f"Working Days: {data.total_working_days} | Days Worked: {data.days_worked}",
    ]
    if data.pan_number:
        info_lines.append(f"PAN: {data.pan_number} | UAN: {data.uan_number or 'N/A'}")
    if data.bank_name:
        info_lines.append(f"Bank: {data.bank_name} | A/C: {data.bank_account or 'N/A'}")
    for line in info_lines:
        _add_body(doc, line, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_paragraph()
    proration = data.days_worked / data.total_working_days if data.total_working_days > 0 else 1
    earnings = [("Basic Salary", data.basic_salary * proration), ("HRA", data.hra * proration)]
    if data.dearness_allowance:
        earnings.append(("DA", data.dearness_allowance * proration))
    if data.conveyance_allowance:
        earnings.append(("Conveyance", data.conveyance_allowance * proration))
    if data.medical_allowance:
        earnings.append(("Medical", data.medical_allowance * proration))
    if data.special_allowance:
        earnings.append(("Special", data.special_allowance * proration))
    if data.other_earnings:
        earnings.append(("Other", data.other_earnings * proration))

    deductions = [("PF", data.pf_deduction)]
    if data.esi_deduction:
        deductions.append(("ESI", data.esi_deduction))
    if data.professional_tax:
        deductions.append(("PT", data.professional_tax))
    if data.income_tax:
        deductions.append(("TDS", data.income_tax))
    if data.other_deductions:
        deductions.append(("Other", data.other_deductions))

    max_rows = max(len(earnings), len(deductions))
    table = doc.add_table(rows=max_rows + 2, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["Earnings", "Amount (₹)", "Deductions", "Amount (₹)"]):
        table.rows[0].cells[i].text = h

    total_e = 0.0
    total_d = 0.0
    for i in range(max_rows):
        if i < len(earnings):
            table.rows[i + 1].cells[0].text = earnings[i][0]
            table.rows[i + 1].cells[1].text = format_inr(earnings[i][1])
            total_e += earnings[i][1]
        if i < len(deductions):
            table.rows[i + 1].cells[2].text = deductions[i][0]
            table.rows[i + 1].cells[3].text = format_inr(deductions[i][1])
            total_d += deductions[i][1]

    table.rows[-1].cells[0].text = "Total Earnings"
    table.rows[-1].cells[1].text = format_inr(total_e)
    table.rows[-1].cells[2].text = "Total Deductions"
    table.rows[-1].cells[3].text = format_inr(total_d)

    doc.add_paragraph()
    _add_body(doc, f"NET PAY: {format_inr(total_e - total_d)}", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _add_subtitle(doc, "This is a system-generated salary slip.")

    _add_footer(doc)
    return _to_bytes(doc)


def generate_legal_notice(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, "LEGAL NOTICE")
    if data.sender_through_advocate:
        _add_subtitle(doc, f"Through: {data.sender_through_advocate}, Advocate")
    _add_body(doc, f"Date: {data.notice_date}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_body(doc, "NOTICE UNDER SECTION 80 CPC / GENERAL LEGAL NOTICE", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    _add_body(doc, f"To,\n{data.recipient_name}\n{data.recipient_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"From,\n{data.sender_name}\n{data.sender_address}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"Subject: {data.subject}", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)

    if data.sender_through_advocate:
        _add_body(doc, f"Under instructions from my client, {data.sender_name}, I serve this legal notice:")
    else:
        _add_body(doc, "I hereby serve this legal notice:")

    _add_section_heading(doc, "FACTS")
    for i, f in enumerate(data.facts):
        _add_body(doc, f"{i + 1}. {f}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "LEGAL GROUNDS")
    for i, g in enumerate(data.legal_grounds):
        _add_body(doc, f"{i + 1}. {g}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_section_heading(doc, "RELIEF SOUGHT")
    for i, r in enumerate(data.relief_sought):
        _add_body(doc, f"{i + 1}. {r}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_body(doc, f"You are called upon to comply within {data.compliance_days} days, failing which legal proceedings will be initiated before the courts in {data.governing_state}.")

    doc.add_paragraph()
    if data.sender_through_advocate:
        _add_body(doc, f"{data.sender_through_advocate}\nAdvocate\nOn behalf of {data.sender_name}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    else:
        _add_body(doc, f"Yours faithfully,\n{data.sender_name}", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_generic_legal_doc(title: str, subtitle: str, sections: list[dict]) -> bytes:
    doc = _create_doc()
    _add_title(doc, title.upper())
    if subtitle:
        _add_subtitle(doc, subtitle)

    for section in sections:
        heading = section.get("heading", "")
        content = section.get("content", "")
        if heading:
            _add_section_heading(doc, heading.upper())
        for para in content.split("\n\n"):
            para = para.strip()
            if para:
                _add_body(doc, para)

    _add_footer(doc)
    return _to_bytes(doc)


def generate_affidavit(data) -> bytes:
    doc = _create_doc()
    _add_title(doc, "AFFIDAVIT")
    _add_subtitle(doc, f"({data.purpose.title()})")

    _add_body(doc, f"I, {data.deponent_name}, aged {data.deponent_age} years, S/o / D/o / W/o {data.deponent_father_name}, Occupation: {data.deponent_occupation}, residing at {data.deponent_address}, do hereby solemnly affirm and state on oath as under:")

    for i, stmt in enumerate(data.statements):
        _add_body(doc, f"{i + 1}. That {stmt}")

    _add_body(doc, f"I declare that the above statements are true and correct to the best of my knowledge and belief.")
    doc.add_paragraph()
    _add_body(doc, f"Place: {data.place}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"Date: {data.date}", alignment=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph()
    _add_body(doc, "DEPONENT", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_body(doc, f"({data.deponent_name})", alignment=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_paragraph()
    _add_body(doc, "VERIFICATION", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _add_body(doc, f"Verified at {data.place} on {data.date} that the contents are true and correct.")

    if data.notary_name:
        doc.add_paragraph()
        _add_body(doc, f"Before me,\n{data.notary_name}\nNotary Public", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    _add_footer(doc)
    return _to_bytes(doc)
